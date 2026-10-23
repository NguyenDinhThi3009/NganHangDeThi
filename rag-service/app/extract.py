"""Trích xuất văn bản từ PDF / DOCX / PPTX. Trả về danh sách (số trang hoặc slide, nội dung)."""
import io
from pathlib import Path

SUPPORTED = {".pdf", ".docx", ".pptx"}


class UnsupportedFormat(Exception):
    pass


def extract_pages(filename: str, data: bytes) -> list[tuple[int, str]]:
    ext = Path(filename).suffix.lower()
    if ext == ".pdf":
        return _pdf(data)
    if ext == ".docx":
        return _docx(data)
    if ext == ".pptx":
        return _pptx(data)
    raise UnsupportedFormat(f"Định dạng {ext or '(không có)'} chưa được hỗ trợ")


def _pdf(data: bytes) -> list[tuple[int, str]]:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    return [(i + 1, page.extract_text() or "") for i, page in enumerate(reader.pages)]


def _docx(data: bytes) -> list[tuple[int, str]]:
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    doc = Document(io.BytesIO(data))
    parts: list[str] = []

    # Đọc theo đúng thứ tự xuất hiện trong tài liệu (đoạn văn xen kẽ bảng)
    for child in doc.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            text = Paragraph(child, doc).text.strip()
            if text:
                parts.append(text)
        elif tag == "tbl":
            for row in Table(child, doc).rows:
                cells: list[str] = []
                prev_tc = None
                for cell in row.cells:
                    if cell._tc is prev_tc:  # ô bị gộp được trả về lặp lại nhiều lần
                        continue
                    prev_tc = cell._tc
                    text = " ".join(cell.text.split())
                    if text and (not cells or text != cells[-1]):
                        cells.append(text)
                if cells:
                    parts.append(" | ".join(cells))

    # DOCX không có khái niệm trang cố định -> gộp thành "trang" 1
    return [(1, "\n".join(parts))]


def _pptx(data: bytes) -> list[tuple[int, str]]:
    from pptx import Presentation

    prs = Presentation(io.BytesIO(data))
    pages: list[tuple[int, str]] = []
    for i, slide in enumerate(prs.slides, start=1):
        texts = []
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                texts.append(shape.text_frame.text.strip())
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
            texts.append(slide.notes_slide.notes_text_frame.text.strip())
        pages.append((i, "\n".join(texts)))
    return pages